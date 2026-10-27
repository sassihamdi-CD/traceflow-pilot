"""Background extraction, branched by file type (Correction 1).

- PDF: file bytes go DIRECTLY to Claude as a native document input.
  No pdfplumber/pypdf pre-extraction. Claude's own PDF reading handles
  layout/tables; the `location` field uses Claude's page reference.
- XLSX: parsed with openpyxl. `location` is the REAL sheet/cell coordinate
  (e.g. "Sheet 1 / F8") taken from openpyxl, never inferred by the model.
  Cell values are flattened into a structured table; Claude only maps that
  table onto field_definitions.
- CSV: streamed with the stdlib csv reader (flat rows, "Row N" locations).
- Unknown types never reach here: the upload endpoint rejects them (415).

Every inserted row uses status='proposed', extracted_by='ai_system'.
Extraction NEVER writes accepted/corrected (spec section 1).
"""
from __future__ import annotations

import asyncio
import base64
import csv
import io
import json
import re
from datetime import date, datetime, time

from anthropic import (
    APIConnectionError,
    APIStatusError,
    AsyncAnthropic,
    InternalServerError,
    RateLimitError,
)
from openpyxl import load_workbook

from app.config import settings
from app.services.extraction_prompts import (
    SYSTEM_PROMPT,
    pdf_task_instruction,
    table_task_instruction,
)


class TruncatedList(list):
    """A plain list that remembers whether a pipeline cap cut it short.

    Behaves exactly like a list (iteration, indexing, len, json.dumps all
    work), so existing call sites are unaffected; new code reads `.truncated`
    to surface the flag on the document row + UI badge.
    """

    def __init__(self, items=(), truncated: bool = False):
        super().__init__(items)
        self.truncated = truncated


# Row caps feeding the truncation flag (mirrors prompt + doc budgets).
TABLE_MAX_ROWS = 500
PARTIES_CAP = 12

# PDF evidence pins. The prompt examples use 'Page 3', but models frequently
# answer 'p. 3' / 'pag. 5'. Normalize instead of silently discarding a valid
# page reference (that mismatch used to null out correct evidence pins).
PDF_LOCATION_RE = re.compile(
    r"^(?:page|p\.?|pag\.?|pagina)\s*(\d+)(?:\s*[-–]\s*(\d+))?$",
    re.IGNORECASE,
)


def table_locations(table: list[dict] | None) -> set[str]:
    """Every location string the model is allowed to echo back."""
    valid: set[str] = set()
    for r in table or []:
        if isinstance(r, dict) and r.get("sheet") and r.get("cell"):
            valid.add(f"{r['sheet']} / {r['cell']}")
    return valid


def coerce_pdf_location(location: object) -> str | None:
    """Normalize a page reference to 'Page N' / 'Page N-M'."""
    if not isinstance(location, str):
        return None
    loc = re.sub(r"\s+", " ", location.strip())
    match = PDF_LOCATION_RE.match(loc)
    if not match:
        return None
    first, last = match.group(1), match.group(2)
    return f"Page {first}-{last}" if last else f"Page {first}"


def coerce_table_location(location: object, valid: set[str]) -> str | None:
    """Keep a sheet/cell pin only when it exists in the parsed table."""
    if not isinstance(location, str):
        return None
    loc = location.strip()
    return loc if loc in valid else None


def validate_item_locations(items: list[dict], ext: str, table: list[dict] | None = None) -> list[dict]:
    """Post-hoc coordinate check on extracted passport values.

    The model is told to copy coordinates verbatim, but it still invents
    them: XLSX/CSV locations must exist in the parsed table, PDF locations
    must match 'Page N'. Unprovable pins are nulled (the VALUE is kept for
    the reviewer) — a fabricated pin is worse than a missing one.
    Never raises; unknown ext keeps locations untouched.
    """
    out: list[dict] = []
    table_exts = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".ods"}
    valid = table_locations(table) if ext in table_exts else set()
    for it in items:
        it = dict(it)
        loc = it.get("location")
        if ext in table_exts:
            it["location"] = coerce_table_location(loc, valid)
        elif ext == ".pdf":
            it["location"] = coerce_pdf_location(loc)
        out.append(it)
    return out


def cell_text(value: object) -> str:
    """Render a spreadsheet cell without corrupting it.

    str() on an openpyxl float turns 12345 into '12345.0' (SKU corruption);
    dates become '2026-10-19 00:00:00' instead of an ISO date; booleans become
    'True'/'False'. Keep the printed value's semantics stable.
    """
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def xlsx_to_table(xlsx_bytes: bytes, max_rows: int = TABLE_MAX_ROWS) -> TruncatedList:
    """Flatten workbook cells to [{sheet, cell, value}]. Coordinates from openpyxl."""
    wb = load_workbook(filename=io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    rows = TruncatedList()
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                        continue
                    if len(rows) >= max_rows:
                        # A (max+1)th non-empty cell exists: the cap bit off data.
                        rows.truncated = True
                        return rows
                    rows.append({"sheet": ws.title, "cell": cell.coordinate, "value": cell_text(cell.value)})
    finally:
        wb.close()
    return rows


def decode_csv_text(csv_bytes: bytes) -> str:
    """Decode UTF-8 (with/without BOM), then Windows-1252, then Latin-1.

    verify_file_content() accepts windows-1252 and latin-1 exports (common
    from Italian Excel), so extraction must accept them too — otherwise an
    accepted upload fails later with a UnicodeDecodeError.
    """
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return csv_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return csv_bytes.decode("latin-1")


def csv_to_table(csv_bytes: bytes, max_rows: int = TABLE_MAX_ROWS) -> TruncatedList:
    """Stream CSV rows into the same [{sheet, cell, value}] shape as XLSX.

    Streaming matters: a 30 MB / 3.6M-cell export used to be materialised as a
    pandas DataFrame (~500 MB RAM per job). The csv reader keeps only the
    decoded text plus the capped cell list.
    """
    text = decode_csv_text(csv_bytes)
    delimiter = ","
    try:
        # Italian/European exports commonly use ';'. A comma-only parser turns
        # the whole row into a single column and loses every field.
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration as exc:
        raise ValueError("csv contains no table data") from exc
    except csv.Error as exc:
        raise ValueError("csv contains no table data") from exc

    headers: list[str] = []
    seen: dict[str, int] = {}
    for idx, name in enumerate(header):
        name = name.strip() if name and name.strip() else f"Unnamed: {idx}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        headers.append(name)

    rows = TruncatedList()
    try:
        for i, record in enumerate(reader, start=2):  # row 1 = header
            for col_idx, value in enumerate(record):
                if value is None or not value.strip():
                    continue
                if len(rows) >= max_rows:
                    rows.truncated = True
                    return rows
                col = headers[col_idx] if col_idx < len(headers) else f"Column {col_idx + 1}"
                rows.append({"sheet": "CSV", "cell": f"Row {i} / {col}", "value": value})
    except csv.Error as exc:
        raise ValueError(f"csv could not be parsed: {exc}") from exc
    return rows


def build_mapping_prompt(field_defs: list[dict], table: list[dict]) -> str:
    allowed = ", ".join(f"{d['field_key']} ({d['label']})" for d in field_defs)
    return table_task_instruction(allowed, json.dumps(table[:500]))


def sanitize_items(items: object, allowed: set[str]) -> list[dict]:
    """Drop anything the model returned that is not a well-formed hit.

    Only well-formed rows survive: known field_key, non-empty value.
    Runs before DB insert; extraction inserts proposed rows only.
    """
    clean: list[dict] = []
    if not isinstance(items, list):
        return clean
    for it in items:
        if not isinstance(it, dict):
            continue
        key = it.get("field_key")
        val = it.get("value")
        if key not in allowed or not isinstance(val, str) or not val.strip():
            continue
        unit = it.get("unit") if isinstance(it.get("unit"), str) else None
        location = it.get("location") if isinstance(it.get("location"), str) else None
        row = {
            "field_key": key,
            "value": val.strip()[:4000],
            "unit": unit.strip()[:50] if unit and unit.strip() else None,
            "location": location.strip()[:200] if location and location.strip() else None,
        }
        # The same field can legitimately repeat on several pages, but an exact
        # duplicate adds no evidence and pollutes the reviewer's queue.
        if row in clean:
            continue
        clean.append(row)
        if len(clean) >= 100:
            break
    return clean


class ExtractionParseError(Exception):
    """Claude's reply contained no parseable JSON payload."""


def parse_extraction_payload(text: str) -> list[dict]:
    """Tolerant parser for the model's JSON reply.

    Handles: raw array, {"items": [...]}, markdown-fenced JSON, and leading /
    trailing prose. Raises ExtractionParseError when no JSON array is found.
    """
    candidate = text.strip()
    if candidate.startswith("```"):
        # strip ```json ... ``` fences
        lines = candidate.splitlines()
        lines = [ln for ln in lines if not ln.strip().startswith("```")]
        candidate = "\n".join(lines).strip()
    try:
        data = json.loads(candidate)
    except json.JSONDecodeError:
        # fall back: first [...] block in the reply
        start, end = candidate.find("["), candidate.rfind("]")
        if start == -1 or end <= start:
            raise ExtractionParseError("no JSON array in model reply")
        try:
            data = json.loads(candidate[start : end + 1])
        except json.JSONDecodeError as e:
            raise ExtractionParseError(f"malformed JSON array: {e}")
    if isinstance(data, dict):
        data = data.get("items", [])
    if not isinstance(data, list):
        raise ExtractionParseError("payload is not a list")
    return data


async def _create_with_retry(client: AsyncAnthropic, **kwargs):
    """Retry transient Anthropic failures (429/5xx/network) with backoff.

    One rate-limit response used to fail a whole document permanently; the
    reviewer then had to re-upload. Three attempts cover normal hiccups.
    """
    delay = 1.0
    last: Exception | None = None
    for attempt in range(3):
        try:
            return await client.messages.create(**kwargs)
        except (RateLimitError, APIConnectionError, InternalServerError) as exc:
            last = exc
        except APIStatusError as exc:
            if exc.status_code not in (429, 500, 502, 503, 504):
                raise
            last = exc
        if attempt < 2:
            await asyncio.sleep(delay)
            delay *= 2
    assert last is not None
    raise last


async def _claude_json_messages(messages: list[dict]) -> list[dict]:
    client = AsyncAnthropic(api_key=settings.anthropic_api_key)
    resp = await _create_with_retry(
        client,
        model=settings.anthropic_model,
        max_tokens=2000,
        system=SYSTEM_PROMPT,
        messages=messages,
    )
    text = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
    return parse_extraction_payload(text)


async def extract_pdf_native(pdf_bytes: bytes, field_defs: list[dict]) -> list[dict]:
    if settings.demo_mode:
        from app.services.demo_extraction import map_pdf

        return map_pdf(pdf_bytes, field_defs)
    allowed = ", ".join(f"{d['field_key']} ({d['label']})" for d in field_defs)
    b64 = base64.standard_b64encode(pdf_bytes).decode("ascii")
    return await _claude_json_messages(
        [
            {
                "role": "user",
                "content": [
                    {
                        "type": "document",
                        "source": {"type": "base64", "media_type": "application/pdf", "data": b64},
                    },
                    {"type": "text", "text": pdf_task_instruction(allowed)},
                ],
            }
        ]
    )


async def extract_table_mapped(table: list[dict], field_defs: list[dict]) -> list[dict]:
    if settings.demo_mode:
        from app.services.demo_extraction import map_table

        return map_table(table, field_defs)
    return await _claude_json_messages(
        [{"role": "user", "content": [{"type": "text", "text": build_mapping_prompt(field_defs, table)}]}]
    )


PARTIES_SYSTEM = (
    "You list the organizations named in a supplier document. "
    "Output STRICT JSON only: an array of objects with exactly the keys "
    "name and role. Role is one of: applicant, vendor, supplier, laboratory, "
    "brand, other. List each distinct organization once with its clearest role. "
    "If none are named, output []. No prose, no markdown."
)


def cap_parties(data: object) -> TruncatedList:
    """Pure parties sanitizer: [{name, role}], capped at PARTIES_CAP.

    Split out of extract_parties_pdf so the 12-cap + truncation flag are
    unit-testable without network. Never raises.
    """
    clean = TruncatedList()
    try:
        if not isinstance(data, list):
            return clean
        kept: list[dict] = []
        for it in data:
            if isinstance(it, dict) and isinstance(it.get("name"), str) and it["name"].strip():
                role = it.get("role") if it.get("role") in (
                    "applicant", "vendor", "supplier", "laboratory", "brand", "other") else "other"
                kept.append({"name": it["name"].strip()[:120], "role": role})
        clean.extend(kept[:PARTIES_CAP])
        clean.truncated = len(kept) > PARTIES_CAP
        return clean
    except Exception:
        return clean


async def extract_parties_pdf(pdf_bytes: bytes) -> list[dict]:
    """Who is on this document? Failure-tolerant: [] on any problem."""
    if settings.demo_mode:
        from app.services.demo_extraction import parties_pdf

        return parties_pdf(pdf_bytes)
    try:
        client = AsyncAnthropic(api_key=settings.anthropic_api_key)
        b64 = base64.standard_b64encode(pdf_bytes).decode("ascii")
        resp = await _create_with_retry(
            client,
            model=settings.anthropic_model,
            max_tokens=400,
            system=PARTIES_SYSTEM,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "document", "source": {
                        "type": "base64", "media_type": "application/pdf", "data": b64}},
                    {"type": "text", "text": "List the organizations named in this document with their roles. JSON array only."},
                ],
            }],
        )
        text = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
        data = parse_extraction_payload(text)
        return cap_parties(data)
    except Exception:
        return []
