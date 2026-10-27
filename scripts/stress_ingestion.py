#!/usr/bin/env python3
"""Hard stress test for TraceFlow ingestion.

Generates synthetic manufacturer documents (Word, Excel, CSV, PDF, text,
images), uploads them concurrently to a running API, waits for the worker to
drain the queue, and validates the stored result.

It is deliberately dependency-free beyond the project's own requirements
(httpx, python-docx, openpyxl, Pillow) and never calls Anthropic directly.

Examples
--------
    # Browse-level load (about 30 documents, small/medium)
    .venv/bin/python scripts/stress_ingestion.py --profile load

    # Large manufacturer packs (200-page PDF, 100k-cell workbook, 10 MB text)
    .venv/bin/python scripts/stress_ingestion.py --profile extreme

    # Type/size/zip-bomb rejection checks (adds one failed worker job)
    .venv/bin/python scripts/stress_ingestion.py --profile boundary

The script prints the product id and file directory so the run can be
inspected or cleaned up afterwards. It does not delete evidence automatically.
"""
from __future__ import annotations

import argparse
import asyncio
import io
import json
import os
import random
import statistics
import string
import sys
import tempfile
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import httpx

DEFAULT_BASE_URL = "http://localhost:8010"
DEFAULT_TOKEN = "demo"
MULTIPART_LIMIT = 4 * 1024 * 1024

FIELD_LINES = [
    ("Product name", "Universal Runner"),
    ("SKU", "STRESS-2026-001"),
    ("Upper material", "Recycled polyester mesh"),
    ("Upper composition", "100% recycled polyester"),
    ("Leather origin", "Italy"),
    ("Sole material", "Rubber"),
    ("REACH compliance", "Compliant with Regulation (EC) 1907/2006"),
]
FIELD_COUNT = len(FIELD_LINES)


@dataclass
class PlannedDoc:
    path: Path
    kind: str
    expected_fields: int
    expect_failure: str | None = None
    expect_status: int | None = None
    note: str = ""
    uploaded_at: float | None = None
    document_id: str | None = None
    upload_status: int | None = None
    upload_seconds: float | None = None
    completed_at: float | None = None


@dataclass
class RunResult:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _rand_text(n: int, rng: random.Random) -> str:
    alphabet = string.ascii_letters + string.digits + " -_/.,()"
    return "".join(rng.choice(alphabet) for _ in range(n))


def _filler_lines(count: int, rng: random.Random) -> list[str]:
    out = []
    for i in range(count):
        out.append(f"Substance {i:05d}: {rng.randint(1, 999)}.{rng.randint(0, 99)} mg/kg")
    return out


# ---------------------------------------------------------------------------
# Synthetic files
# ---------------------------------------------------------------------------
def make_text(path: Path, lines: int, rng: random.Random, pad: int = 0) -> None:
    with path.open("w", encoding="utf-8") as fh:
        for label, value in FIELD_LINES:
            fh.write(f"{label}: {value}\n")
        for line in _filler_lines(lines, rng):
            fh.write(line + "\n")
        if pad:
            fh.write("X" * pad)


def make_csv(path: Path, rows: int, cols: int, rng: random.Random) -> None:
    import csv

    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["label", "value"] + [f"col{j}" for j in range(2, cols)])
        for label, value in FIELD_LINES:
            w.writerow([label, value] + ["" for _ in range(2, cols)])
        for i in range(rows):
            w.writerow([f"row{i}", f"value {i}"] + [
                f"{rng.randint(0, 99999)}" for _ in range(2, cols)
            ])


def make_xlsx(path: Path, rows: int, cols: int, rng: random.Random) -> None:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Declaration")
    ws.append(["label", "value"] + [f"col{j}" for j in range(2, cols)])
    for label, value in FIELD_LINES:
        ws.append([label, value] + ["" for _ in range(2, cols)])
    for i in range(rows):
        ws.append([f"row{i}", f"value {i}"] + [
            rng.randint(0, 99999) for _ in range(2, cols)
        ])
    wb.save(path)


def make_docx(path: Path, paragraphs: int, rng: random.Random) -> None:
    from docx import Document

    doc = Document()
    for label, value in FIELD_LINES:
        doc.add_paragraph(f"{label}: {value}")
    for line in _filler_lines(paragraphs, rng):
        doc.add_paragraph(line)
    doc.save(path)


def _pdf_escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: int, lines_per_page: int, rng: random.Random) -> None:
    """Minimal valid PDF with a standard Helvetica text stream."""

    def page_lines(index: int) -> list[str]:
        if index == 0:
            lines = [f"{label}: {value}" for label, value in FIELD_LINES]
            lines += _filler_lines(max(lines_per_page - len(lines), 0), rng)
            return lines
        return _filler_lines(lines_per_page, rng)

    objects: list[bytes] = [b""]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(pages))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>".encode())
    objects.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding >>"
    )
    for i in range(pages):
        body = (
            "BT\n/F1 9 Tf\n40 800 Td\n11 TL\n"
            + "".join(f"({_pdf_escape(line)}) Tj T*\n" for line in page_lines(i))
            + "ET\n"
        )
        raw = body.encode("latin-1", "replace")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(
            f"<< /Length {len(raw)} >>\nstream\n".encode() + raw + b"endstream"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for i, obj in enumerate(objects[1:], start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects)}\n0000000000 65535 f \n".encode())
    for off in offsets[1:]:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects)} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(out.getvalue())


def make_png(path: Path, width: int, height: int, rng: random.Random) -> None:
    from PIL import Image

    img = Image.new("RGB", (width, height), (240, 240, 240))
    # A small pattern keeps the PNG from compressing into a tiny file.
    pixels = img.load()
    for y in range(0, height, 7):
        for x in range(0, width, 11):
            pixels[x, y] = (rng.randint(0, 255), rng.randint(0, 255), rng.randint(0, 255))
    img.save(path, format="PNG")


def make_zip_bomb(path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"\0" * (20 * 1024 * 1024))


def make_fake_pdf(path: Path, size_mb: int) -> None:
    with path.open("wb") as fh:
        fh.write(b"%PDF-1.7\n")
        remaining = size_mb * 1024 * 1024 - 9
        chunk = b"0" * (1024 * 1024)
        while remaining > 0:
            take = min(remaining, len(chunk))
            fh.write(chunk[:take])
            remaining -= take


# ---------------------------------------------------------------------------
# Planning
# ---------------------------------------------------------------------------
def plan_documents(profile: str, root: Path, rng: random.Random) -> list[PlannedDoc]:
    docs: list[PlannedDoc] = []

    def add(name: str, kind: str, expected: int, **kw) -> None:
        docs.append(PlannedDoc(root / name, kind, expected, **kw))

    if profile in {"smoke", "load", "extreme", "boundary"}:
        for i in range(3):
            name = f"declaration-{i}.docx"
            make_docx(root / name, 120, rng)
            add(name, "docx", FIELD_COUNT)

        for i in range(2):
            name = f"bom-{i}.csv"
            make_csv(root / name, 800, 5, rng)
            add(name, "csv", FIELD_COUNT)

        for i in range(2):
            name = f"spec-{i}.xlsx"
            make_xlsx(root / name, 400, 5, rng)
            add(name, "xlsx", FIELD_COUNT)

        name = "notes.txt"
        make_text(root / name, 500, rng)
        add(name, "text", FIELD_COUNT)

        name = "pack-30p.pdf"
        make_pdf(root / name, 30, 120, rng)
        add(name, "pdf", FIELD_COUNT)

        name = "label.png"
        make_png(root / name, 1200, 1800, rng)
        add(name, "image", 0)

    if profile in {"load", "extreme"}:
        for i in range(6):
            name = f"declaration-many-{i}.docx"
            make_docx(root / name, 600, rng)
            add(name, "docx", FIELD_COUNT)

        for i in range(4):
            name = f"bom-many-{i}.csv"
            make_csv(root / name, 4000, 6, rng)
            add(name, "csv", FIELD_COUNT)

        for i in range(4):
            name = f"spec-many-{i}.xlsx"
            make_xlsx(root / name, 2500, 6, rng)
            add(name, "xlsx", FIELD_COUNT)

        for i in range(3):
            name = f"long-notes-{i}.txt"
            make_text(root / name, 30000, rng)
            add(name, "text", FIELD_COUNT)

        for i in range(2):
            name = f"pack-100p-{i}.pdf"
            make_pdf(root / name, 100, 180, rng)
            add(name, "pdf", FIELD_COUNT)

    if profile == "extreme":
        name = "pack-200p.pdf"
        make_pdf(root / name, 200, 300, rng)
        add(name, "pdf", FIELD_COUNT, note="text cap + item cap")

        name = "huge-workbook.xlsx"
        make_xlsx(root / name, 20000, 5, rng)
        add(name, "xlsx", FIELD_COUNT, note="50k cell cap")

        name = "huge-table.csv"
        make_csv(root / name, 30000, 6, rng)
        add(name, "csv", FIELD_COUNT, note="50k cell cap")

        name = "huge-docx.docx"
        make_docx(root / name, 4000, rng)
        add(name, "docx", FIELD_COUNT, note="4000 paragraphs")

        name = "huge-text.txt"
        make_text(root / name, 300000, rng)
        add(name, "text", FIELD_COUNT, note="4 MB text cap")

    if profile == "large":
        name = "large-text-12mb.txt"
        make_text(root / name, 100000, rng, pad=10 * 1024 * 1024)
        add(name, "text", FIELD_COUNT, note="~12 MB text, 4 MB cap")

        name = "large-table-30mb.csv"
        make_csv(root / name, 300000, 12, rng)
        add(name, "csv", FIELD_COUNT, note="~25-30 MB CSV / 3.6M cells")

        name = "large-spec-1m-cells.xlsx"
        make_xlsx(root / name, 100000, 10, rng)
        add(name, "xlsx", FIELD_COUNT, note="1M cells, 50k cell cap")

    if profile == "boundary":
        name = "zip-bomb.docx"
        make_zip_bomb(root / name)
        add(name, "docx", 0, expect_status=415, expect_failure="zip_bomb")

        name = "too-big.pdf"
        make_fake_pdf(root / name, 65)
        add(name, "pdf", 0, expect_status=413, expect_failure="file_too_large")

        name = "corrupt.pdf"
        (root / name).write_bytes(b"%PDF-1.7\nthis is not a real pdf")
        add(name, "pdf", 0, note="worker must fail cleanly", expect_failure="pdf_corrupt")

    return docs


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------
class StressClient:
    def __init__(self, base_url: str, token: str, timeout: float):
        self.base_url = base_url.rstrip("/")
        self.token = token
        self.timeout = timeout
        self.headers = {"Authorization": f"Bearer {token}"}

    async def create_product(self, sku: str) -> str:
        async with httpx.AsyncClient(timeout=self.timeout) as client:
            r = await client.post(
                f"{self.base_url}/api/products",
                headers=self.headers,
                json={
                    "sku": sku,
                    "name": f"Stress {sku}",
                    "category": "footwear",
                    "manufacturer_name": "Stress Supplier Srl",
                },
            )
            r.raise_for_status()
            return r.json()["id"]

    async def upload(self, product_id: str, doc: PlannedDoc) -> tuple[int, dict]:
        content = doc.path.read_bytes()
        started = time.monotonic()
        async with httpx.AsyncClient(timeout=self.timeout) as client:
            r = await client.post(
                f"{self.base_url}/api/products/{product_id}/documents",
                headers=self.headers,
                files={"file": (doc.path.name, content)},
            )
        doc.upload_seconds = time.monotonic() - started
        doc.upload_status = r.status_code
        try:
            data = r.json()
        except ValueError:
            data = {"raw": r.text[:300]}
        return r.status_code, data

    async def detail(self, product_id: str) -> dict:
        async with httpx.AsyncClient(timeout=self.timeout) as client:
            r = await client.get(
                f"{self.base_url}/api/products/{product_id}", headers=self.headers
            )
            r.raise_for_status()
            return r.json()


def _pct(values: list[float], p: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    idx = min(len(ordered) - 1, int(round((p / 100) * (len(ordered) - 1))))
    return ordered[idx]


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------
async def run(
    profile: str,
    base_url: str,
    token: str,
    concurrency: int,
    timeout: float,
    out_dir: Path,
) -> RunResult:
    result = RunResult()
    rng = random.Random(20261026)
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Generating {profile} documents in {out_dir} ...")
    docs = plan_documents(profile, out_dir, rng)
    total_bytes = sum(d.path.stat().st_size for d in docs)
    print(f"{len(docs)} documents, {total_bytes / 1024 / 1024:.1f} MiB on disk")

    client = StressClient(base_url, token, timeout)
    sku = f"STRESS-{profile.upper()}-{int(time.time())}"
    product_id = await client.create_product(sku)
    print(f"Product: {product_id} ({sku})")

    semaphore = asyncio.Semaphore(concurrency)

    async def upload_one(doc: PlannedDoc) -> None:
        async with semaphore:
            try:
                status, data = await client.upload(product_id, doc)
                doc.document_id = data.get("document_id")
                if doc.expect_status is not None:
                    if status != doc.expect_status:
                        result.errors.append(
                            f"{doc.path.name}: expected HTTP {doc.expect_status}, got {status} {data}"
                        )
                elif status != 200:
                    result.errors.append(f"{doc.path.name}: upload HTTP {status} {data}")
            except Exception as exc:  # noqa: BLE001
                result.errors.append(f"{doc.path.name}: upload raised {type(exc).__name__}: {exc}")

    started = time.monotonic()
    await asyncio.gather(*(upload_one(d) for d in docs))
    upload_elapsed = time.monotonic() - started
    accepted = [d for d in docs if d.document_id]
    print(
        f"Uploads finished in {upload_elapsed:.1f}s: "
        f"{len(accepted)} accepted, {len(docs) - len(accepted)} rejected/skipped"
    )

    if not accepted:
        _report_uploads(docs, result)
        return result

    print("Waiting for the worker queue to drain ...")
    deadline = time.monotonic() + timeout
    last_state: dict[str, str] = {}
    while time.monotonic() < deadline:
        detail = await client.detail(product_id)
        by_id = {d["id"]: d for d in detail.get("documents", [])}
        pending = 0
        for doc in accepted:
            row = by_id.get(doc.document_id or "")
            if not row:
                pending += 1
                continue
            state = row.get("extraction_status")
            last_state[doc.document_id or ""] = state or "missing"
            if state in {"done", "failed"}:
                if doc.completed_at is None:
                    doc.completed_at = time.monotonic()
            else:
                pending += 1
        if pending == 0:
            break
        await asyncio.sleep(1.0)

    drain_elapsed = time.monotonic() - started - upload_elapsed
    detail = await client.detail(product_id)
    _report_uploads(docs, result)
    _validate(docs, detail, result)
    _print_summary(docs, detail, upload_elapsed, drain_elapsed, result, product_id, out_dir)
    return result


def _report_uploads(docs: list[PlannedDoc], result: RunResult) -> None:
    latencies = [d.upload_seconds for d in docs if d.upload_seconds is not None]
    if latencies:
        print(
            "Upload latency: "
            f"p50={_pct(latencies, 50):.2f}s p95={_pct(latencies, 95):.2f}s "
            f"max={max(latencies):.2f}s"
        )


def _validate(docs: list[PlannedDoc], detail: dict, result: RunResult) -> None:
    documents = {d["id"]: d for d in detail.get("documents", [])}
    fields = detail.get("fields", [])

    for doc in docs:
        if doc.expect_status is not None:
            continue
        row = documents.get(doc.document_id or "")
        if row is None:
            result.errors.append(f"{doc.path.name}: document missing after upload")
            continue
        status = row.get("extraction_status")
        if status == "extracting":
            result.errors.append(f"{doc.path.name}: still extracting after timeout")
        elif doc.expect_failure:
            if status != "failed":
                result.errors.append(f"{doc.path.name}: expected failure, got {status}")
            elif doc.expect_failure != "any" and doc.expect_failure not in str(row.get("extraction_error") or ""):
                result.errors.append(
                    f"{doc.path.name}: failure did not mention {doc.expect_failure}: "
                    f"{row.get('extraction_error')}"
                )
            continue
        else:
            if status != "done":
                result.errors.append(
                    f"{doc.path.name}: expected done, got {status} ({row.get('extraction_error')})"
                )
            elif row.get("extracted_count") != doc.expected_fields:
                result.errors.append(
                    f"{doc.path.name}: expected {doc.expected_fields} fields, "
                    f"got {row.get('extracted_count')}"
                )
        if row.get("items_count", 0) > 2000:
            result.errors.append(f"{doc.path.name}: items_count={row.get('items_count')} > 2000")
        if row.get("extraction_version") != "grounded-v1":
            result.errors.append(
                f"{doc.path.name}: extraction_version={row.get('extraction_version')}"
            )
        if row.get("truncated"):
            result.warnings.append(f"{doc.path.name}: truncated={row.get('truncated')} (expected for large docs)")

    for field in fields:
        verified = [
            v for v in field.get("values", [])
            if v.get("status") in {"accepted", "corrected"}
        ]
        if len(verified) > 1:
            result.errors.append(
                f"field {field.get('field_key')}: {len(verified)} verified values"
            )


def _print_summary(
    docs: list[PlannedDoc],
    detail: dict,
    upload_elapsed: float,
    drain_elapsed: float,
    result: RunResult,
    product_id: str,
    out_dir: Path,
) -> None:
    documents = {d["id"]: d for d in detail.get("documents", [])}
    done = sum(1 for d in docs if d.document_id and documents.get(d.document_id, {}).get("extraction_status") == "done")
    failed = sum(1 for d in docs if d.document_id and documents.get(d.document_id, {}).get("extraction_status") == "failed")
    rejected = sum(1 for d in docs if not d.document_id)
    items = sum(int(documents.get(d.document_id, {}).get("items_count") or 0) for d in docs if d.document_id)
    proposals = sum(
        1
        for f in detail.get("fields", [])
        for v in f.get("values", [])
        if v.get("status") == "proposed"
    )
    print("")
    print("=== STRESS RESULT ===")
    print(f"product_id        {product_id}")
    print(f"files             {len(docs)}")
    print(f"done / failed     {done} / {failed}")
    print(f"rejected/skipped  {rejected}")
    print(f"document items    {items}")
    print(f"proposed fields   {proposals}")
    print(f"upload wall time  {upload_elapsed:.1f}s")
    print(f"drain wall time   {drain_elapsed:.1f}s")
    print(f"files kept at     {out_dir}")
    if result.warnings:
        print(f"warnings          {len(result.warnings)}")
        for w in result.warnings[:15]:
            print(f"  - {w}")
    if result.errors:
        print(f"ERRORS            {len(result.errors)}")
        for e in result.errors[:40]:
            print(f"  - {e}")
    else:
        print("ERRORS            0")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", choices=["smoke", "load", "extreme", "large", "boundary"], default="load")
    parser.add_argument("--base-url", default=os.environ.get("TRACEFLOW_API", DEFAULT_BASE_URL))
    parser.add_argument("--token", default=os.environ.get("TRACEFLOW_TOKEN", DEFAULT_TOKEN))
    parser.add_argument("--concurrency", type=int, default=8, help="parallel uploads")
    parser.add_argument("--timeout", type=float, default=900.0, help="overall drain timeout")
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args(argv)
    out_dir = args.out_dir or Path(tempfile.mkdtemp(prefix="traceflow-stress-"))
    result = asyncio.run(run(
        args.profile, args.base_url, args.token,
        args.concurrency, args.timeout, out_dir,
    ))
    print(json.dumps({"errors": result.errors[:50], "warnings": result.warnings[:50]}, indent=2))
    return 1 if result.errors else 0


if __name__ == "__main__":
    sys.exit(main())
